"""
service/document_extractor.py

Document text extraction for survey input.

Extracts plain text from uploaded BRD/FRD/PRD files so the survey
pipeline can use the content as requirement context.

Supports:
  - Plain text (.txt)
  - PDF via PyMuPDF (.pdf)
  - Word documents via python-docx (.docx, .doc)

Returns None on failure rather than raising — a document extraction
failure must not abort the whole survey.
"""

import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def _extract_pdf_page(page) -> str:
    """
    Extract text from a single PDF page, trying structured table extraction first.

    A requirements table in a PDF (ID | description | acceptance criteria) is
    fragmented into disconnected lines by `page.get_text()`. PyMuPDF 1.23+
    provides `page.find_tables()` which reconstructs the row/column structure
    so the requirement ID stays on the same line as its description.

    Strategy:
      1. Attempt `page.find_tables()` (PyMuPDF >= 1.23). Each table row is
         joined as tab-separated cells and appended as a line. A blank line
         separates tables.
      2. Always append the full `page.get_text()` as well — prose requirements
         may appear alongside or instead of tables, and we must not lose them.
      3. If `find_tables` is unavailable (older PyMuPDF), fall back to
         `page.get_text()` only. The fallback is logged at DEBUG level so it
         is visible when diagnosing ID-extraction failures on PDFs.

    Args:
        page: A PyMuPDF page object.

    Returns:
        Extracted text for this page as a single string.
    """
    lines: list[str] = []
    table_rows_found = 0

    try:
        tables = page.find_tables()
        for table in tables:
            for row in table.extract():
                cells = [str(c or "").strip() for c in row]
                if any(cells):
                    lines.append("\t".join(cells))
                    table_rows_found += 1
            lines.append("")  # blank line between tables
        if table_rows_found:
            logger.debug(
                "_extract_pdf_page: found %d table row(s) via find_tables()",
                table_rows_found,
            )
    except AttributeError:
        # PyMuPDF < 1.23 does not have find_tables().
        logger.debug(
            "_extract_pdf_page: find_tables() unavailable (PyMuPDF < 1.23) "
            "— falling back to get_text() only"
        )

    # Always include the full page text. Tables may be mixed with prose, and
    # get_text() is the authoritative source for non-tabular content.
    lines.append(page.get_text())
    return "\n".join(lines)


def extract_text_from_document(path: str) -> Optional[str]:
    """
    Extract plain text from an uploaded document file.

    Args:
        path: Path to the file on disk.

    Returns:
        Extracted text as a string, or None if extraction failed.
    """
    path_obj = Path(path)
    suffix = path_obj.suffix.lower()

    try:
        if suffix == ".txt":
            return path_obj.read_text(encoding="utf-8")

        if suffix == ".pdf":
            import fitz  # PyMuPDF
            doc = fitz.open(path)
            text = "\n".join(_extract_pdf_page(page) for page in doc)
            logger.info("Extracted %d chars from PDF: %s", len(text), path_obj.name)
            return text

        if suffix in (".docx", ".doc"):
            # 0-D: Reject old Word 97-2003 binary .doc files before attempting
            # to open them with python-docx (which only handles the newer .docx
            # ZIP-based format).  The OLE2 Compound File signature is the first
            # four bytes of every .doc / .xls / .ppt from Office 97-2003.
            if suffix == ".doc":
                with open(path, "rb") as _f:
                    _magic = _f.read(4)
                if _magic == b"\xd0\xcf\x11\xe0":
                    logger.warning(
                        "'%s' is a Word 97-2003 (.doc) binary file. "
                        "python-docx cannot read this format. "
                        "Please re-save the document as .docx and re-upload.",
                        path_obj.name,
                    )
                    return None

            import docx
            doc = docx.Document(path)

            # 0-A: Extract paragraph text (existing behaviour)
            para_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

            # 0-A: Also extract table rows — Baxter BRDs store requirement IDs
            # like BXU_RPT1_GEN5_V2.0 in Word tables, not in paragraphs. Without
            # this loop, every table cell is silently dropped and all requirement
            # IDs come back as null.
            #
            # Merged-cell fix: python-docx returns the SAME cell object multiple
            # times in row.cells — once per column a merged cell spans. Joining
            # them naively duplicates the text (e.g. "BXU_RPT1\tBXU_RPT1").
            # We track the underlying XML element identity (_tc) and skip any
            # cell whose element we've already seen in this row.
            table_lines: list[str] = []
            for table in doc.tables:
                for row in table.rows:
                    seen_tc: set[int] = set()
                    cells: list[str] = []
                    for cell in row.cells:
                        tc_id = id(cell._tc)   # unique per physical XML cell element
                        if tc_id in seen_tc:
                            continue           # merged cell — already included
                        seen_tc.add(tc_id)
                        text = cell.text.strip()
                        if text:
                            cells.append(text)
                    if cells:
                        table_lines.append("\t".join(cells))
                table_lines.append("")  # blank line between tables
            table_text = "\n".join(table_lines).strip()

            text = "\n\n".join(part for part in [para_text, table_text] if part)
            logger.info(
                "Extracted %d chars from DOCX (%d paragraphs, %d tables): %s",
                len(text), len(doc.paragraphs), len(doc.tables), path_obj.name,
            )
            return text

    except Exception as e:
        logger.warning(
            "Could not extract text from %s (%s): %s",
            path_obj.name, suffix, e,
        )

    return None
